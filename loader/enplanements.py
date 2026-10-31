"""FAA boardings loader: the `enplanements` table (docs/data-sources.md B, docs/schema.md).

What it is:   Downloads the FAA boardings (enplanements) Excel files and loads them into
              the `enplanements` table: boardings and FAA hub size per airport and year.
Why we need:  Boardings give airport size (the size filter, the hub size) and the
              short-term growth from 2024 to 2025.
Used by:      loader/build.py. tests: test_step02_enplanements.py,
              test_step03_code_renames.py, test_step05_build.py (through build).
Uses:         the downloaded FAA .xlsx files; the FAA -> IATA code map made by
              loader/airports.py (passed in by build); writes the enplanements table.
              loader/common.py (CODE_RENAMES, LoadReport).
Main parts:   FAA_URL - the download address.
              YEAR_LABEL - finds the year in a column header such as `CY 25 Enplanements`.
              download_faa(years, dest_dir) - download the file for each year; return paths.
              load_enplanements(xlsx_paths, faa_to_iata, con) - replace the enplanements
                table from the .xlsx files.
Design:       docs/data-sources.md B; docs/schema.md `enplanements`; docs/decisions.md
              7.6; docs/implementation-plan.md P7.

Each FAA file gives two years (e.g. `CY 25 Enplanements`, `CY 24 Enplanements`).
Downloaded files are untrusted: columns are found by header name, values go to
SQL as parameters, never as text.
"""

import re
from pathlib import Path

import httpx
from openpyxl import load_workbook

from loader.common import CODE_RENAMES, LoadReport

FAA_URL = (
    "https://www.faa.gov/airports/planning_capacity/passenger_allcargo_stats/"
    "passenger/arp-cy{year}-commercial-service-enplanements.xlsx"
)

YEAR_LABEL = re.compile(r"CY\s*(\d{2})\b")


def download_faa(years, dest_dir) -> list[Path]:
    """Download the FAA boardings file for each year into dest_dir; return the paths."""
    dest = Path(dest_dir)
    dest.mkdir(parents=True, exist_ok=True)
    paths = []
    for year in years:
        url = FAA_URL.format(year=int(year))
        path = dest / url.rsplit("/", 1)[1]
        response = httpx.get(url, timeout=120, follow_redirects=True)
        response.raise_for_status()
        path.write_bytes(response.content)
        paths.append(path)
    return paths


def _read_file(path):
    """Return ({year: column index}, header name -> index, data rows) of the first sheet."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    columns = {name: i for i, name in enumerate(header)}
    year_cols = {}
    for i, name in enumerate(header):
        match = YEAR_LABEL.search(name)
        if match:
            year_cols[2000 + int(match.group(1))] = i
    width = len(header)
    data = [tuple(r) + (None,) * (width - len(r)) for r in rows[1:]]  # short rows padded
    return year_cols, columns, data


def load_enplanements(xlsx_paths, faa_to_iata: dict, con) -> LoadReport:
    """Replace the contents of the `enplanements` table from the FAA .xlsx files.

    When two files give the same (airport, year), the file with the newer year
    labels wins (P7). Old codes are renamed first (CODE_RENAMES, decisions 7.6),
    then FAA codes are translated only through `faa_to_iata`;
    untranslated codes are skipped and listed in the warnings.
    """
    files = [_read_file(p) for p in xlsx_paths]
    files.sort(key=lambda f: max(f[0]), reverse=True)  # newest first

    values = {}  # (iata, year) -> (enplanements, hub_type)
    warnings = []
    warned = set()
    for year_cols, columns, rows in files:
        locid, hub, state, name = (
            columns["Locid"], columns["Hub"], columns["ST"], columns["Airport Name"]
        )
        for row in rows:
            faa = row[locid]
            numbers = {
                year: int(row[col])
                for year, col in year_cols.items()
                if isinstance(row[col], (int, float)) and not isinstance(row[col], bool)
            }
            if not isinstance(faa, str) or not faa.strip() or not numbers:
                continue  # blank or summary row
            faa = CODE_RENAMES.get(faa.strip(), faa.strip())
            iata = faa_to_iata.get(faa)
            if iata is None:
                if faa not in warned:
                    warned.add(faa)
                    warnings.append(
                        f"FAA code {faa} ({row[state]}, {row[name]}) has no IATA code; not loaded"
                    )
                continue
            hub_type = row[hub]
            hub_type = None if hub_type is None or str(hub_type).strip() == "None" else str(hub_type).strip()
            for year, n in numbers.items():
                values.setdefault((iata, year), (n, hub_type))

    con.execute("DELETE FROM enplanements")
    if values:
        con.executemany(
            "INSERT INTO enplanements (iata, year, enplanements, hub_type) VALUES (?, ?, ?, ?)",
            [[iata, year, n, hub_type] for (iata, year), (n, hub_type) in values.items()],
        )
    rows = con.execute("SELECT count(*) FROM enplanements").fetchone()[0]
    return LoadReport(table="enplanements", rows=rows, warnings=warnings)
